import io, datetime
from typing import Dict, Any

DISCLAIMER_TEXT = "AI-generated information is for understanding purposes and is not a substitute for professional legal advice."

def export_as_txt(document, analysis: Dict[str, Any]) -> bytes:
    """Generate a clean, structured plain-text report of the analysis."""
    lines = []
    lines.append("=" * 72)
    lines.append(" AI DOC — Intelligent Document Analysis & AI Assistant")
    lines.append("=" * 72)
    lines.append(f"Document Name : {document.original_filename}")
    lines.append(f"File Type     : {document.file_type} ({document.file_size} bytes)")
    lines.append(f"Document Type : {analysis.get('document_type', 'General')}")
    lines.append(f"Export Date   : {datetime.datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S UTC')}")
    lines.append(f"Engine        : {analysis.get('analysis_engine', 'AI DOC Engine')}")
    lines.append("-" * 72)
    lines.append("")

    # Summary
    lines.append("## EXECUTIVE SUMMARY")
    lines.append(analysis.get('summary', 'No summary available.'))
    lines.append("")

    # Main Purpose
    if analysis.get('main_purpose'):
        lines.append("## MAIN PURPOSE")
        lines.append(analysis.get('main_purpose'))
        lines.append("")

    # Key Points
    key_points = analysis.get('key_points', [])
    if key_points:
        lines.append("## KEY POINTS")
        for idx, kp in enumerate(key_points, 1):
            lines.append(f"  {idx}. {kp}")
        lines.append("")

    # Important Dates
    dates = analysis.get('important_dates', [])
    if dates:
        lines.append("## IMPORTANT DATES")
        for d in dates:
            dtype = d.get('type', 'Date')
            dstr = d.get('date', 'N/A')
            devt = d.get('event', '')
            lines.append(f"  * [{dtype}] {dstr} — {devt}")
        lines.append("")

    # Parties / Entities
    parties = analysis.get('parties_or_entities', [])
    if parties:
        lines.append("## PARTIES / ENTITIES")
        for p in parties:
            pname = p.get('name', 'Unknown')
            prole = p.get('role', 'Participant')
            ptype = p.get('type', 'Entity')
            lines.append(f"  * {pname} ({prole} / {ptype})")
        lines.append("")

    # Obligations
    obligations = analysis.get('obligations', [])
    if obligations:
        lines.append("## OBLIGATIONS & RESPONSIBILITIES")
        for o in obligations:
            party = o.get('party', 'Party')
            resp = o.get('responsibility', '')
            lines.append(f"  * [{party}]: {resp}")
        lines.append("")

    # Risks / Warnings
    risks = analysis.get('risks_or_warnings', [])
    if risks:
        lines.append("## RISKS & WARNINGS")
        for r in risks:
            title = r.get('title', 'Notice')
            sev = r.get('severity', 'Attention')
            desc = r.get('description', '')
            lines.append(f"  * [{sev.upper()}] {title}: {desc}")
        lines.append("")

    # Action Items
    actions = analysis.get('action_items', [])
    if actions:
        lines.append("## RECOMMENDED ACTION ITEMS")
        for idx, a in enumerate(actions, 1):
            lines.append(f"  [ ] {idx}. {a}")
        lines.append("")

    # Simplified Explanation
    simp = analysis.get('simplified_explanation', '')
    if simp:
        lines.append("## SIMPLIFIED EXPLANATION")
        lines.append(simp)
        lines.append("")

    # Key Terms
    terms = analysis.get('key_terms', [])
    if terms:
        lines.append("## KEY TERMS GLOSSARY")
        for t in terms:
            term = t.get('term', '')
            exp = t.get('explanation', '')
            lines.append(f"  * {term}: {exp}")
        lines.append("")

    # Legal Disclaimer
    lines.append("=" * 72)
    lines.append(f"DISCLAIMER: {DISCLAIMER_TEXT}")
    lines.append("=" * 72)

    return "\n".join(lines).encode('utf-8')

def export_as_pdf(document, analysis: Dict[str, Any]) -> bytes:
    """Generate a clean PDF report using ReportLab."""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=40, leftMargin=40,
            topMargin=40, bottomMargin=40
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Heading1'],
            fontSize=20,
            leading=24,
            textColor=colors.HexColor('#1e293b'),
            spaceAfter=4
        )
        subtitle_style = ParagraphStyle(
            'DocSubtitle',
            parent=styles['Normal'],
            fontSize=11,
            leading=14,
            textColor=colors.HexColor('#64748b'),
            spaceAfter=12
        )
        h2_style = ParagraphStyle(
            'H2',
            parent=styles['Heading2'],
            fontSize=13,
            leading=16,
            textColor=colors.HexColor('#0f172a'),
            spaceBefore=12,
            spaceAfter=6
        )
        body_style = ParagraphStyle(
            'Body',
            parent=styles['Normal'],
            fontSize=9.5,
            leading=13.5,
            textColor=colors.HexColor('#334155'),
            spaceAfter=6
        )
        bullet_style = ParagraphStyle(
            'Bullet',
            parent=body_style,
            leftIndent=15,
            spaceAfter=4
        )
        disclaimer_style = ParagraphStyle(
            'Disclaimer',
            parent=styles['Italic'],
            fontSize=8,
            leading=11,
            textColor=colors.HexColor('#64748b'),
            alignment=1
        )

        elements = []

        # Header
        elements.append(Paragraph("AI DOC", title_style))
        elements.append(Paragraph("Intelligent Document Analysis Report", subtitle_style))
        elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#3b82f6'), spaceAfter=10))

        # Meta Table
        meta_data = [
            [Paragraph("<b>Document:</b>", body_style), Paragraph(document.original_filename, body_style),
             Paragraph("<b>Type:</b>", body_style), Paragraph(analysis.get('document_type', 'General'), body_style)],
            [Paragraph("<b>File Size:</b>", body_style), Paragraph(f"{document.file_size:,} bytes", body_style),
             Paragraph("<b>Date:</b>", body_style), Paragraph(datetime.datetime.utcnow().strftime('%Y-%m-%d'), body_style)]
        ]
        meta_table = Table(meta_data, colWidths=[65, 200, 50, 200])
        meta_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#f8fafc')),
            ('PADDING', (0,0), (-1,-1), 4),
            ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor('#e2e8f0')),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e2e8f0')),
        ]))
        elements.append(meta_table)
        elements.append(Spacer(1, 10))

        # Summary
        elements.append(Paragraph("Executive Summary", h2_style))
        summary_text = analysis.get('summary', 'No summary available.')
        for p in summary_text.split('\n\n'):
            elements.append(Paragraph(p.strip(), body_style))

        # Key Points
        key_points = analysis.get('key_points', [])
        if key_points:
            elements.append(Paragraph("Key Points", h2_style))
            for idx, kp in enumerate(key_points, 1):
                elements.append(Paragraph(f"• <b>{idx}.</b> {kp}", bullet_style))

        # Important Dates
        dates = analysis.get('important_dates', [])
        if dates:
            elements.append(Paragraph("Important Dates & Milestones", h2_style))
            date_rows = [[Paragraph("<b>Type</b>", body_style), Paragraph("<b>Date</b>", body_style), Paragraph("<b>Event / Context</b>", body_style)]]
            for d in dates:
                date_rows.append([
                    Paragraph(d.get('type', 'Date'), body_style),
                    Paragraph(d.get('date', 'N/A'), body_style),
                    Paragraph(d.get('event', ''), body_style)
                ])
            date_table = Table(date_rows, colWidths=[90, 95, 330])
            date_table.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#f1f5f9')),
                ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
                ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e2e8f0')),
                ('PADDING', (0,0), (-1,-1), 4),
            ]))
            elements.append(date_table)
            elements.append(Spacer(1, 8))

        # Obligations
        obligations = analysis.get('obligations', [])
        if obligations:
            elements.append(Paragraph("Obligations & Responsibilities", h2_style))
            for o in obligations:
                party = o.get('party', 'Party')
                resp = o.get('responsibility', '')
                elements.append(Paragraph(f"• <b>[{party}]:</b> {resp}", bullet_style))

        # Risks
        risks = analysis.get('risks_or_warnings', [])
        if risks:
            elements.append(Paragraph("Risks & Significant Conditions", h2_style))
            for r in risks:
                title = r.get('title', 'Notice')
                sev = r.get('severity', 'Attention')
                desc = r.get('description', '')
                elements.append(Paragraph(f"• <b>[{sev.upper()}] {title}:</b> {desc}", bullet_style))

        # Action Items
        actions = analysis.get('action_items', [])
        if actions:
            elements.append(Paragraph("Action Items", h2_style))
            for idx, a in enumerate(actions, 1):
                elements.append(Paragraph(f"[  ] <b>{idx}.</b> {a}", bullet_style))

        # Key Terms
        terms = analysis.get('key_terms', [])
        if terms:
            elements.append(Paragraph("Key Terms & Glossary", h2_style))
            for t in terms:
                term = t.get('term', '')
                exp = t.get('explanation', '')
                elements.append(Paragraph(f"• <b>{term}:</b> {exp}", bullet_style))

        # Footer Disclaimer
        elements.append(Spacer(1, 15))
        elements.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#94a3b8'), spaceAfter=8))
        elements.append(Paragraph(f"DISCLAIMER: {DISCLAIMER_TEXT}", disclaimer_style))

        doc.build(elements)
        buffer.seek(0)
        return buffer.getvalue()
    except Exception as e:
        print(f"[Export Service] PDF generation error: {e}. Falling back to TXT bytes.")
        return export_as_txt(document, analysis)

def export_as_docx(document, analysis: Dict[str, Any]) -> bytes:
    """Generate a Microsoft Word DOCX report."""
    try:
        import docx
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = docx.Document()
        
        # Title
        p_title = doc.add_paragraph()
        run_title = p_title.add_run("AI DOC")
        run_title.font.size = Pt(22)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(30, 41, 59)

        p_sub = doc.add_paragraph("Intelligent Document Analysis Report")
        p_sub.runs[0].font.size = Pt(12)
        p_sub.runs[0].font.color.rgb = RGBColor(100, 116, 139)

        # Metadata
        doc.add_paragraph(f"Document Name: {document.original_filename}")
        doc.add_paragraph(f"Document Type: {analysis.get('document_type', 'General')}")
        doc.add_paragraph(f"Analysis Date: {datetime.datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}")
        doc.add_paragraph(f"Analysis Engine: {analysis.get('analysis_engine', 'AI DOC')}")
        doc.add_paragraph()

        # Summary
        doc.add_heading("Executive Summary", level=1)
        doc.add_paragraph(analysis.get('summary', 'No summary available.'))

        # Main Purpose
        if analysis.get('main_purpose'):
            doc.add_heading("Main Purpose", level=1)
            doc.add_paragraph(analysis.get('main_purpose'))

        # Key Points
        key_points = analysis.get('key_points', [])
        if key_points:
            doc.add_heading("Key Points", level=1)
            for kp in key_points:
                doc.add_paragraph(kp, style='List Bullet')

        # Important Dates
        dates = analysis.get('important_dates', [])
        if dates:
            doc.add_heading("Important Dates", level=1)
            table = doc.add_table(rows=1, cols=3)
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = 'Type'
            hdr_cells[1].text = 'Date'
            hdr_cells[2].text = 'Event / Context'
            for d in dates:
                row_cells = table.add_row().cells
                row_cells[0].text = d.get('type', 'Date')
                row_cells[1].text = d.get('date', 'N/A')
                row_cells[2].text = d.get('event', '')

        # Obligations
        obligations = analysis.get('obligations', [])
        if obligations:
            doc.add_heading("Obligations & Responsibilities", level=1)
            for o in obligations:
                party = o.get('party', 'Party')
                resp = o.get('responsibility', '')
                doc.add_paragraph(f"[{party}] {resp}", style='List Bullet')

        # Risks
        risks = analysis.get('risks_or_warnings', [])
        if risks:
            doc.add_heading("Risks & Significant Conditions", level=1)
            for r in risks:
                title = r.get('title', 'Notice')
                sev = r.get('severity', 'Attention')
                desc = r.get('description', '')
                doc.add_paragraph(f"[{sev.upper()}] {title}: {desc}", style='List Bullet')

        # Action Items
        actions = analysis.get('action_items', [])
        if actions:
            doc.add_heading("Action Items", level=1)
            for a in actions:
                doc.add_paragraph(f"[ ] {a}", style='List Bullet')

        # Key Terms
        terms = analysis.get('key_terms', [])
        if terms:
            doc.add_heading("Key Terms", level=1)
            for t in terms:
                term = t.get('term', '')
                exp = t.get('explanation', '')
                doc.add_paragraph(f"{term}: {exp}", style='List Bullet')

        # Disclaimer
        doc.add_paragraph()
        p_disc = doc.add_paragraph(f"DISCLAIMER: {DISCLAIMER_TEXT}")
        p_disc.runs[0].font.size = Pt(8.5)
        p_disc.runs[0].font.italic = True
        p_disc.runs[0].font.color.rgb = RGBColor(100, 116, 139)

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
    except Exception as e:
        print(f"[Export Service] DOCX generation error: {e}. Falling back to TXT bytes.")
        return export_as_txt(document, analysis)
