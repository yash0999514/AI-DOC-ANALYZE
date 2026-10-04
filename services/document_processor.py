import os, re
from typing import Dict, Any, Tuple
from services.ocr_service import ocr_image, ocr_pdf_pages

ALLOWED_EXTENSIONS = {'pdf', 'docx', 'txt', 'png', 'jpg', 'jpeg'}

def allowed_file(filename: str) -> bool:
    if not filename or '.' not in filename:
        return False
    ext = filename.rsplit('.', 1)[1].lower()
    return ext in ALLOWED_EXTENSIONS

def detect_file_type(filename: str) -> str:
    if not filename or '.' not in filename:
        return 'UNKNOWN'
    ext = filename.rsplit('.', 1)[1].lower()
    if ext == 'pdf':
        return 'PDF'
    elif ext == 'docx':
        return 'DOCX'
    elif ext == 'txt':
        return 'TXT'
    elif ext in {'png', 'jpg', 'jpeg'}:
        return 'IMAGE'
    return ext.upper()

def clean_text(text: str) -> str:
    if not text:
        return ""
    # Remove null bytes
    text = text.replace('\x00', '')
    # Normalize CRLF and CR to LF
    text = text.replace('\r\n', '\n').replace('\r', '\n')
    # Collapse 3 or more newlines into 2
    text = re.sub(r'\n{3,}', '\n\n', text)
    # Strip leading and trailing whitespace
    return text.strip()

def extract_from_pdf(file_path: str) -> Tuple[str, int, bool]:
    """Extract text from PDF natively, falling back to OCR if scanned."""
    native_text_parts = []
    page_count = 0
    is_scanned = False

    try:
        import pypdf
        reader = pypdf.PdfReader(file_path)
        page_count = len(reader.pages)
        for i, page in enumerate(reader.pages):
            try:
                page_text = page.extract_text() or ''
                if page_text.strip():
                    native_text_parts.append(f"--- [Page {i + 1}] ---\n{page_text.strip()}")
            except Exception as pe:
                print(f"[DocProcessor] Native extract failed on page {i + 1}: {pe}")
    except Exception as e:
        print(f"[DocProcessor] pypdf reader error: {e}")

    full_native_text = "\n\n".join(native_text_parts).strip()
    total_non_ws = len(re.sub(r'\s+', '', full_native_text))

    # Detect if PDF is scanned (little or no text found)
    if total_non_ws < 50 or (page_count > 0 and total_non_ws / page_count < 15):
        print(f"[DocProcessor] PDF {file_path} appears to be scanned ({total_non_ws} chars). Invoking OCR...")
        ocr_text = ocr_pdf_pages(file_path, max_pages=15)
        if ocr_text:
            return clean_text(ocr_text), page_count, True
        # If OCR returned empty, return whatever native text we got
        return clean_text(full_native_text), page_count, True

    return clean_text(full_native_text), page_count, False

def extract_from_docx(file_path: str) -> Tuple[str, int, bool]:
    """Extract text from DOCX paragraphs and tables."""
    import docx
    doc = docx.Document(file_path)
    text_parts = []

    for para in doc.paragraphs:
        txt = para.text.strip()
        if txt:
            text_parts.append(txt)

    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                text_parts.append(" | ".join(row_cells))

    full_text = "\n\n".join(text_parts).strip()
    return clean_text(full_text), 1, False

def extract_from_txt(file_path: str) -> Tuple[str, int, bool]:
    """Extract text from TXT with robust encoding fallbacks."""
    encodings = ['utf-8', 'utf-8-sig', 'latin-1', 'cp1252', 'iso-8859-1']
    for enc in encodings:
        try:
            with open(file_path, 'r', encoding=enc) as f:
                content = f.read()
                return clean_text(content), 1, False
        except (UnicodeDecodeError, UnicodeError):
            continue

    # Fallback with ignore errors
    with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
        content = f.read()
        return clean_text(content), 1, False

def extract_from_image(file_path: str) -> Tuple[str, int, bool]:
    """Extract text from PNG/JPG image via OCR."""
    text = ocr_image(file_path)
    return clean_text(text), 1, True

def process_document(file_path: str, filename: str) -> Dict[str, Any]:
    """Process an uploaded document file: identify type, extract text, run OCR if needed."""
    if not os.path.isfile(file_path):
        return {
            'success': False,
            'error': f"File does not exist: {file_path}"
        }

    file_type = detect_file_type(filename)
    try:
        if file_type == 'PDF':
            extracted_text, page_count, is_scanned = extract_from_pdf(file_path)
        elif file_type == 'DOCX':
            extracted_text, page_count, is_scanned = extract_from_docx(file_path)
        elif file_type == 'TXT':
            extracted_text, page_count, is_scanned = extract_from_txt(file_path)
        elif file_type == 'IMAGE':
            extracted_text, page_count, is_scanned = extract_from_image(file_path)
        else:
            return {
                'success': False,
                'error': f"Unsupported file type: {file_type}. Please upload PDF, DOCX, TXT, JPG, or PNG."
            }

        extracted_text = clean_text(extracted_text)
        if not extracted_text:
            return {
                'success': False,
                'error': "We couldn't extract readable text from this document. It may be empty or corrupted."
            }

        words = extracted_text.split()
        return {
            'success': True,
            'file_type': file_type,
            'extracted_text': extracted_text,
            'is_scanned': is_scanned,
            'page_count': page_count,
            'char_count': len(extracted_text),
            'word_count': len(words)
        }
    except Exception as e:
        return {
            'success': False,
            'error': f"Error processing {file_type} document: {str(e)}"
        }
